"""Génération du classeur Excel de synthèse (openpyxl)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from carburant.models import PleinCarburant
from dashboard import services
from depenses.models import AutreDepense
from entretien.models import Entretien
from poclain.models import ActiviteEngin
from portechar.models import TrajetPorteChar
from voyages.models import Voyage

_ENTETE = Font(bold=True, color="FFFFFF")
_FOND = PatternFill("solid", fgColor="0D6EFD")
_TITRE = Font(bold=True, size=14)


def _ecrire_tableau(ws, entetes, lignes, depart=1):
    for j, e in enumerate(entetes, start=1):
        c = ws.cell(row=depart, column=j, value=e)
        c.font = _ENTETE
        c.fill = _FOND
    for i, ligne in enumerate(lignes, start=depart + 1):
        for j, val in enumerate(ligne, start=1):
            ws.cell(row=i, column=j, value=val)
    for j in range(1, len(entetes) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 18


def construire(annee=None):
    wb = Workbook()

    # --- Synthèse ---
    ws = wb.active
    ws.title = "Synthèse"
    ws["A1"] = "Gestion de flotte — Synthèse"
    ws["A1"].font = _TITRE
    ws["A2"] = f"Année : {annee or 'Toutes'}"
    t = services.totaux(annee)
    _ecrire_tableau(ws, ["Indicateur", "Valeur"], [
        ["DÉPENSES", None],
        ["Carburant véhicules", float(t["carburant_vehicules"])],
        ["Carburant Poclain", float(t["carburant_poclain"])],
        ["Entretien", float(t["montant_entretien"])],
        ["Autres dépenses", float(t["montant_autres"])],
        ["Total dépenses", float(t["depenses_totales"])],
        ["RECETTES", None],
        ["Voyages camions", float(t["recettes_voyages"])],
        ["Porte-char", float(t["recettes_portechar"])],
        ["Chargements Poclain", float(t["recettes_poclain"])],
        ["Total recettes", float(t["recettes_totales"])],
        ["RÉSULTAT (recettes − dépenses)", float(t["resultat"])],
        ["ACTIVITÉ", None],
        ["Carburant consommé (litres)", float(t["litres_carburant"])],
        ["Nombre de voyages", t["nb_voyages"]],
        ["Nombre de chargements Poclain", t["nb_chargements"]],
        ["Nombre de véhicules", t["nb_vehicules"]],
    ], depart=4)

    # --- Analyse par véhicule ---
    wsa = wb.create_sheet("Analyse véhicules")
    lignes = services.analyse_vehicules(annee)
    _ecrire_tableau(wsa, [
        "Véhicule", "Distance (km)", "Litres", "L/100km",
        "Carburant", "Entretien", "Coût/km carbur.", "Coût/km entret.",
        "Coût total/km", "Voyages",
    ], [[
        l["vehicule"], l["distance"], l["litres"], l["l_100km"],
        l["montant_carburant"], l["montant_entretien"],
        l["cout_carburant_km"], l["cout_entretien_km"],
        l["cout_total_km"], l["nb_voyages"],
    ] for l in lignes])

    # --- Feuilles de données brutes ---
    def filtrer(qs):
        # .all() garantit un QuerySet itérable même si on reçoit un Manager.
        return qs.filter(date__year=annee) if annee else qs.all()

    _ecrire_tableau(
        wb.create_sheet("Entretien"),
        ["Date", "Véhicule", "Garage", "Nature", "Montant TTC"],
        [[e.date, str(e.vehicule), e.garage, e.nature_operations, float(e.montant_ttc)]
         for e in filtrer(Entretien.objects.select_related("vehicule"))],
    )
    _ecrire_tableau(
        wb.create_sheet("Carburant"),
        ["Date", "Véhicule", "Litres", "Prix/L", "Montant"],
        [[p.date, str(p.vehicule), float(p.litres), float(p.prix_litre), float(p.montant)]
         for p in filtrer(PleinCarburant.objects.select_related("vehicule"))],
    )
    _ecrire_tableau(
        wb.create_sheet("Voyages"),
        ["Date", "Véhicule", "Nb voyages", "Montant"],
        [[v.date, str(v.vehicule), v.nb_voyage, float(v.montant)]
         for v in filtrer(Voyage.objects.select_related("vehicule"))],
    )
    _ecrire_tableau(
        wb.create_sheet("Poclain"),
        ["Date", "Engin", "Nb chargements", "Montant chargement", "Montant carburant"],
        [[a.date, str(a.engin), a.nb_chargement, float(a.montant_chargement), float(a.montant_carburant)]
         for a in filtrer(ActiviteEngin.objects.select_related("engin"))],
    )
    _ecrire_tableau(
        wb.create_sheet("Porte-char"),
        ["Date", "Engin", "Départ", "Arrivée", "Montant payé"],
        [[tr.date, str(tr.engin or ""), tr.point_depart, tr.point_arrivee, float(tr.montant_paye)]
         for tr in filtrer(TrajetPorteChar.objects.select_related("engin"))],
    )
    _ecrire_tableau(
        wb.create_sheet("Autres dépenses"),
        ["Date", "Désignation", "PU", "Quantité", "Montant"],
        [[d.date, d.designation, float(d.prix_unitaire), float(d.quantite), float(d.montant)]
         for d in filtrer(AutreDepense.objects)],
    )
    return wb
