"""
Importe le classeur « Gestion_Stock » (feuilles Inventaire + Mouvements).

Usage :
    python manage.py import_stock "chemin/Gestion_Stock Recent (1).xlsx"

Ré-exécutable : les articles sont repérés par leur code (pas de doublon).
Les mouvements ne sont PAS dédupliqués automatiquement ; n'importez qu'une fois
le journal, sinon les stocks seront comptés deux fois.
"""
import datetime
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError

from stock.models import Article, MouvementStock


def dec(v):
    if v in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(v).replace(" ", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def to_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if v in (None, ""):
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            continue
    return None


class Command(BaseCommand):
    help = "Importe l'inventaire et les mouvements depuis le classeur de stock."

    def add_arguments(self, parser):
        parser.add_argument("fichier")
        parser.add_argument("--sans-mouvements", action="store_true",
                            help="N'importer que les articles (ignorer le journal).")

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("openpyxl est requis (pip install openpyxl).")
        try:
            wb = openpyxl.load_workbook(options["fichier"], data_only=True)
        except FileNotFoundError:
            raise CommandError("Fichier introuvable.")
        except Exception as exc:
            raise CommandError(f"Lecture impossible : {exc}")

        self._import_articles(wb)
        if not options["sans_mouvements"]:
            self._import_mouvements(wb)
        self.stdout.write(self.style.SUCCESS("Import du stock terminé."))

    def _import_articles(self, wb):
        if "Inventaire" not in wb.sheetnames:
            self.stdout.write("Feuille 'Inventaire' absente.")
            return
        ws = wb["Inventaire"]
        n = maj = 0
        # En-têtes en ligne 4, données à partir de la ligne 5.
        for row in ws.iter_rows(min_row=5, values_only=True):
            code = (str(row[0]).strip() if row[0] else "")
            designation = str(row[1] or "").strip()
            # On ignore les lignes incomplètes (code sans désignation).
            if not code or not designation:
                continue
            specification = str(row[2] or "").strip()
            categorie = str(row[3] or "").strip()
            stock_initial = dec(row[4])
            seuil_mini = dec(row[8]) if len(row) > 8 else Decimal("0")
            observation = str(row[10] or "").strip() if len(row) > 10 else ""

            a = Article.all_objects.filter(code=code).first()
            if a:
                a.designation = designation or a.designation
                a.specification = specification
                a.categorie = categorie
                a.stock_initial = stock_initial
                a.seuil_mini = seuil_mini
                a.observation = observation
                a.is_deleted = False
                a.save()
                maj += 1
            else:
                Article(
                    code=code, designation=designation, specification=specification,
                    categorie=categorie, stock_initial=stock_initial,
                    seuil_mini=seuil_mini, observation=observation,
                ).save()
                n += 1
        self.stdout.write(f"  Articles : {n} créé(s), {maj} mis à jour.")

    def _import_mouvements(self, wb):
        if "Mouvements" not in wb.sheetnames:
            self.stdout.write("Feuille 'Mouvements' absente.")
            return
        ws = wb["Mouvements"]
        n = ignores = 0
        for row in ws.iter_rows(min_row=5, values_only=True):
            code = (str(row[1]).strip() if len(row) > 1 and row[1] else "")
            if not code:
                continue
            art = Article.all_objects.filter(code=code).first()
            if not art:
                ignores += 1
                continue
            type_txt = str(row[3] or "").strip().lower() if len(row) > 3 else ""
            type_code = MouvementStock.ENTREE if type_txt.startswith("entr") else MouvementStock.SORTIE
            MouvementStock(
                article=art,
                date=to_date(row[0]) or datetime.date.today(),
                type=type_code,
                quantite=dec(row[4]) if len(row) > 4 else Decimal("0"),
                motif=str(row[5] or "").strip() if len(row) > 5 else "",
                responsable=str(row[6] or "").strip() if len(row) > 6 else "",
            ).save()
            n += 1
        self.stdout.write(f"  Mouvements : {n} importé(s)" + (f", {ignores} ignoré(s) (code inconnu)." if ignores else "."))
